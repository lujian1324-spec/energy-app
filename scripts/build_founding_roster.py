#!/usr/bin/env python3
"""
Regenerate src/data/foundingMemberRoster.ts from the VIP export.

    python scripts/build_founding_roster.py <workbook.xlsx>

The roster ships inside the app, and the app ships to anyone who wants it, so
the addresses are NOT written into it. Each one is stored as the SHA-256 of its
normalised (lowercased, trimmed) form, and the app hashes the address of
whoever signs in and looks that up. The lookup answers exactly as it would with
plain addresses; what changes is that unpacking the bundle no longer hands
someone the customer list. An attacker can still test a specific address they
already know — hashing cannot prevent that — but 223 members' addresses can no
longer be READ OUT of the app, which is the exposure that matters.

Keep the spreadsheet as the source of truth and re-run this when it changes;
the generated file is not meant to be edited by hand.

Reads the "Email Lookup" sheet, which carries one row per unique address
(aliases included) against the member number it belongs to.
"""
import hashlib
import sys
from pathlib import Path

try:
    import openpyxl
except ImportError:
    sys.exit("openpyxl is needed to read the workbook:  pip install openpyxl")

SHEET = "Email Lookup"

# Addresses that are not on the sheet but must be on the roster: the team's own
# accounts, carrying numbers deliberately outside the sheet's 1..N range so they
# can never collide with a real member's place in the order. Kept here rather
# than in the workbook so the export stays a clean record of who actually
# qualified, and so regenerating never drops them.
EXTRA: dict[str, int] = {
    "jason@sierro.us": 666,
    "benson8191@gmail.com": 999,
}
OUT = Path(__file__).resolve().parent.parent / "src" / "data" / "foundingMemberRoster.ts"


def normalise(email: str) -> str:
    """Must match normalizeEmail() in foundingMembers.ts exactly."""
    return str(email).strip().lower()


def digest(email: str) -> str:
    return hashlib.sha256(normalise(email).encode("utf-8")).hexdigest()


def read_rows(path: Path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if SHEET not in wb.sheetnames:
        sys.exit(f'no "{SHEET}" sheet in {path.name}; found {wb.sheetnames}')
    rows = list(wb[SHEET].iter_rows(values_only=True))
    try:
        header = next(i for i, r in enumerate(rows) if r and r[0] == "Founding Member No.")
    except StopIteration:
        sys.exit(f'no "Founding Member No." header row in the {SHEET} sheet')
    out = []
    for r in rows[header + 1:]:
        if not r or r[0] is None or not r[1]:
            continue
        out.append((normalise(r[1]), int(r[0])))
    return out


def main() -> None:
    if len(sys.argv) != 2:
        sys.exit(__doc__)
    src = Path(sys.argv[1])
    if not src.exists():
        sys.exit(f"no such file: {src}")

    pairs = read_rows(src)
    if not pairs:
        sys.exit("the sheet produced no rows — refusing to write an empty roster")

    sheet_members = sorted({n for _, n in pairs})
    for email, number in EXTRA.items():
        if number in sheet_members:
            sys.exit(f"extra {email} wants #{number}, which the sheet already assigns")
        pairs.append((normalise(email), number))

    # An address appearing twice would make the member it maps to depend on row
    # order, so it is a mistake in the sheet rather than something to resolve here.
    seen: dict[str, int] = {}
    for email, number in pairs:
        if email in seen and seen[email] != number:
            sys.exit(f"{email} is on the sheet against both #{seen[email]} and #{number}")
        seen[email] = number

    members = sorted({n for _, n in pairs})
    by_hash = {digest(e): n for e, n in pairs}
    if len(by_hash) != len(seen):
        sys.exit("two addresses hashed to the same value — this should be impossible")

    lines = [
        "/**",
        " * GENERATED — do not edit by hand.",
        " *",
        " *   python scripts/build_founding_roster.py <workbook.xlsx>",
        " *",
        " * SHA-256 of each founding member's normalised address -> their member",
        " * number. The addresses themselves are deliberately not here: this file",
        " * ships inside the app, and a readable list would hand anyone who unpacks",
        " * the bundle every VIP customer's email. See foundingMembers.ts for how it",
        " * is read, and the script above for how it is built.",
        " *",
        f" * {len(members)} members across {len(by_hash)} addresses, from {src.name}"
        f" (#{sheet_members[0]}-#{sheet_members[-1]}) plus"
        f" {len(EXTRA)} team accounts (#{', #'.join(str(n) for n in sorted(EXTRA.values()))}).",
        " */",
        "",
        "export const FOUNDING_MEMBER_HASHES: Readonly<Record<string, number>> = {",
    ]
    # Sorted by hash, so the diff between two exports shows only real changes
    # rather than the sheet's row order.
    for h in sorted(by_hash):
        lines.append(f"  '{h}': {by_hash[h]},")
    lines.append("}")
    lines.append("")

    OUT.write_text("\n".join(lines), encoding="utf-8")
    print(f"{OUT.relative_to(OUT.parents[2])}: {len(by_hash)} addresses, {len(members)} members")


if __name__ == "__main__":
    main()
